"""Builds the impact model workbook. Every result is a formula; inputs are blue and documented."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Supplier_Follow-up_Agent_Impact_Model.xlsx")
wb = Workbook()

ARIAL = "Arial"
F_BASE = Font(name=ARIAL, size=10)
F_BOLD = Font(name=ARIAL, size=10, bold=True)
F_TITLE = Font(name=ARIAL, size=15, bold=True, color="172B43")
F_SUB = Font(name=ARIAL, size=10, italic=True, color="627187")
F_H = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_IN = Font(name=ARIAL, size=10, color="0000FF")            # hardcoded inputs
F_LINK = Font(name=ARIAL, size=10, color="008000")          # links to another sheet
F_LINKB = Font(name=ARIAL, size=10, color="008000", bold=True)
F_CALC = Font(name=ARIAL, size=10, color="000000")
F_CALCB = Font(name=ARIAL, size=10, color="000000", bold=True)
F_NOTE = Font(name=ARIAL, size=9, color="42506B")
FILL_H = PatternFill("solid", fgColor="142338")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_TOT = PatternFill("solid", fgColor="EEF1F7")
FILL_SEC = PatternFill("solid", fgColor="E6ECFF")
THIN = Side(style="thin", color="DFE5EE")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

INR = '[>=10000000]"₹"##\\,##\\,##\\,##0;[>=100000]"₹"##\\,##\\,##0;"₹"#,##0'
INR2 = '"₹"#,##0.00'
LAKH = '"₹"#,##0.0" L"'
PCT = '0.0%'
PCT0 = '0%'
NUM = '#,##0'
X = '0.0"x"'
MON = '0.0" months"'
USD = '"$"#,##0.00'

def style_range(ws, rng, font=None, fill=None, border=None, align=None, fmt=None):
    for row in ws[rng]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if border: c.border = border
            if align: c.alignment = align
            if fmt: c.number_format = fmt

def header(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font = F_H; c.fill = FILL_H; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BOX

def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def title(ws, text, sub):
    ws["A1"] = text; ws["A1"].font = F_TITLE
    ws["A2"] = sub; ws["A2"].font = F_SUB
    ws.sheet_view.showGridLines = False

# ---------------------------------------------------------------- Inputs
inp = wb.active; inp.title = "Inputs"
title(inp, "Inputs", "Blue cells are assumptions: change them and every other sheet updates. None of these has been measured yet.")
header(inp, 4, ["#", "Input", "Unit", "Cautious", "Likely", "Strong", "Why this value, and how it gets measured"])
INPUTS = [
    # key, label, unit, cautious, likely, strong, fmt, note
    ("F", "Flagged supplier-side credit per month", "₹", 1000000, 2500000, 5000000, INR,
     "Assumption for a mid-size buyer with several GSTINs. The demo company has ₹23.4 L flagged across 16 suppliers. Measured from reco results in shadow mode."),
    ("L0", "Share of flagged credit lost for good today", "%", 0.03, 0.06, 0.10, PCT,
     "Hypothesis H4: credit lost when last year's invoices miss the 30 Nov deadline (Section 16(4)) or suppliers never file (Rule 37A). No public data. The first number shadow mode measures; below 3% the case weakens."),
    ("L1", "Share lost with the agent", "%", 0.02, 0.02, 0.01, PCT,
     "Some suppliers never file whatever we do (closed, in distress), so loss does not reach zero."),
    ("FT", "Share fixed within one filing cycle today", "%", 0.60, 0.55, 0.50, PCT,
     "Assumption: manual chasing with promises tracked in spreadsheets. Measured in shadow mode."),
    ("FW", "Share fixed within one filing cycle with the agent", "%", 0.70, 0.80, 0.85, PCT,
     "First message the day the reco job flags, reminders timed before the 11th, broken promises caught on the 12th. Pilot KR1 targets 70%."),
    ("DS", "Days sooner for credit that moves into the first cycle", "days", 30, 60, 75, NUM,
     "A missed cycle costs about a month; credit that slips usually slips two cycles."),
    ("CC", "Cost of capital", "% a year", 0.10, 0.12, 0.14, PCT,
     "Working-capital rate for a mid-size Indian manufacturer (assumption)."),
    ("S", "Suppliers followed up per month", "count", 40, 60, 100, NUM,
     "Assumption. The demo has 16 suppliers open on one day."),
    ("MT", "Minutes per supplier per cycle today", "minutes", 30, 40, 50, NUM,
     "Reading replies, chasing, updating the tracker (assumption; hypothesis H1 asks this in interviews)."),
    ("MW", "Minutes per supplier per cycle with the agent", "minutes", 18, 12, 10, NUM,
     "Checking readings and approving drafts. Pilot KR3 targets under 15."),
    ("H", "Loaded cost of one GST team hour", "₹", 500, 700, 900, INR,
     "Assumption."),
    ("P", "Price per customer per year", "₹", 200000, 200000, 200000, INR,
     "Placeholder for a paid add-on. Pricing is an open question (per GSTIN, per supplier, or a share of credit recovered)."),
    ("N", "Customers adopting (portfolio view)", "count", 60, 150, 300, NUM,
     "Assumption for illustration, not Masters India data."),
]
R = {}
for i, (k, label, unit, lo, mid, hi, fmt, note) in enumerate(INPUTS):
    r = 5 + i; R[k] = r
    inp.cell(row=r, column=1, value=i + 1).font = F_BASE
    inp.cell(row=r, column=2, value=label).font = F_BOLD
    inp.cell(row=r, column=3, value=unit).font = F_BASE
    for j, v in enumerate((lo, mid, hi)):
        c = inp.cell(row=r, column=4 + j, value=v); c.font = F_IN; c.number_format = fmt
        if j == 1: c.fill = FILL_KEY
    inp.cell(row=r, column=7, value=note).font = F_NOTE
    for col in range(1, 8):
        inp.cell(row=r, column=col).border = BOX; inp.cell(row=r, column=col).alignment = WRAP
last_in = 5 + len(INPUTS) - 1
inp.cell(row=last_in + 2, column=2, value="Scenarios: Cautious sets every input at the low-value end at once, Strong at the high end. The Sensitivity sheet moves one input at a time.").font = F_NOTE
inp.cell(row=last_in + 3, column=2, value="Yellow = the likely case used in the document and slides. Blue = typed input; black = formula; green = taken from another sheet.").font = F_NOTE
widths(inp, [4, 44, 10, 13, 13, 13, 90])
inp.freeze_panes = "D5"

def ref(k, col):  # col: 'D' cautious, 'E' likely, 'F' strong
    return f"Inputs!${col}${R[k]}"

SC = [("Cautious", "D"), ("Likely", "E"), ("Strong", "F")]

# ---------------------------------------------------------------- AI unit economics
ai = wb.create_sheet("AI unit economics")
title(ai, "AI unit economics", "What the AI costs per reply, per supplier per cycle and per customer per year, at list prices on 11 Oct 2026.")
header(ai, 4, ["Model", "Input (US$ per million tokens)", "Output (US$ per million tokens)", "Role"])
PRICES = [("Claude Haiku 5.5", 0.10, 0.50, "Runs every job (likely case)"),
          ("Claude Sonnet 5.5", 2.0, 10.0, "Retry after a failed check (likely case)"),
          ("Claude Haiku 4.5", 1.0, 5.0, "Older quick model: the high-cost case"),
          ("Claude Sonnet 4.5", 3.0, 15.0, "Retry in the high-cost case")]
for i, (m, pi, po, role) in enumerate(PRICES):
    r = 5 + i
    ai.cell(row=r, column=1, value=m).font = F_BOLD
    c = ai.cell(row=r, column=2, value=pi); c.font = F_IN; c.number_format = USD
    c = ai.cell(row=r, column=3, value=po); c.font = F_IN; c.number_format = USD
    ai.cell(row=r, column=4, value=role).font = F_BASE
    for col in range(1, 5): ai.cell(row=r, column=col).border = BOX
ai["A9"] = "Source: Anthropic pricing page (platform.claude.com/docs/en/about-claude/pricing), opened 11 Oct 2026; Haiku 5.5 price for prompts up to 100,000 tokens."; ai["A9"].font = F_NOTE
ai["A11"] = "Rupees per US dollar"; ai["A11"].font = F_BOLD
ai["B11"] = 97; ai["B11"].font = F_IN; ai["B11"].number_format = '"₹"0.00'; ai["B11"].fill = FILL_KEY
ai["C11"] = "Mid-market rate ₹96.64 on 11 Oct 2026 (Upstox currency converter), rounded up."; ai["C11"].font = F_NOTE
ai["A12"] = "Share of calls retried on the larger model"; ai["A12"].font = F_BOLD
ai["B12"] = 0.10; ai["B12"].font = F_IN; ai["B12"].number_format = PCT0; ai["B12"].fill = FILL_KEY
ai["C12"] = "Assumption. No reading failed its checks in the 40-reply blind test, so 1 in 10 is cautious."; ai["C12"].font = F_NOTE

header(ai, 14, ["Job", "Input tokens per call", "Output tokens per call", "Calls per supplier per cycle", "Input tokens per cycle", "Output tokens per cycle"])
JOBS = [("Read a reply", 1430, 270, 2), ("Draft a message", 500, 190, 3), ("Summarise the supplier", 1150, 140, 3), ("Answer a question", 1000, 100, 1)]
for i, (j, ti, to, n) in enumerate(JOBS):
    r = 15 + i
    ai.cell(row=r, column=1, value=j).font = F_BOLD
    for col, v in ((2, ti), (3, to), (4, n)):
        c = ai.cell(row=r, column=col, value=v); c.font = F_IN; c.number_format = NUM
    c = ai.cell(row=r, column=5, value=f"=B{r}*D{r}"); c.font = F_CALC; c.number_format = NUM
    c = ai.cell(row=r, column=6, value=f"=C{r}*D{r}"); c.font = F_CALC; c.number_format = NUM
    for col in range(1, 7): ai.cell(row=r, column=col).border = BOX
ai["A19"] = "Total per supplier per cycle"; ai["A19"].font = F_BOLD
ai["E19"] = "=SUM(E15:E18)"; ai["F19"] = "=SUM(F15:F18)"
for a in ("E19", "F19"): ai[a].font = F_CALCB; ai[a].number_format = NUM; ai[a].fill = FILL_TOT
ai["A20"] = "Tokens estimated as characters ÷ 3 from the prototype's real prompts and recorded outputs (mixed English and Indian languages). Calls per cycle: first message, reminder and follow-up; two replies; three summary refreshes; one question."; ai["A20"].font = F_NOTE

header(ai, 22, ["Result", "Likely (Haiku 5.5, retries on Sonnet 5.5)", "High cost (Haiku 4.5, retries on Sonnet 4.5)", "How it is calculated"])
rows = [
    ("Cost of reading one reply, no retry (₹)", "=(B15*B5+C15*C5)/1000000*$B$11", "=(B15*B7+C15*C7)/1000000*$B$11", "Tokens × price ÷ 1,000,000 × rupees per dollar", INR2),
    ("Cost per supplier per cycle (₹)", "=((E19*B5+F19*C5)+$B$12*(E19*B6+F19*C6))/1000000*$B$11", "=((E19*B7+F19*C7)+$B$12*(E19*B8+F19*C8))/1000000*$B$11", "All calls on the quick model, plus the retried share again on the larger model", INR2),
    ("Suppliers followed up per year (likely)", f"={ref('S','E')}*12", f"={ref('S','E')}*12", "Suppliers per month × 12", NUM),
    ("AI cost per customer per year (₹)", "=B24*B25", "=C24*C25", "Cost per supplier-cycle × suppliers per year", INR),
    ("AI cost as a share of the price", f"=B26/{ref('P','E')}", f"=C26/{ref('P','E')}", "AI cost ÷ price per customer per year", '0.00%'),
]
for i, (lab, f1, f2, how, fmt) in enumerate(rows):
    r = 23 + i
    ai.cell(row=r, column=1, value=lab).font = F_BOLD
    for col, f in ((2, f1), (3, f2)):
        c = ai.cell(row=r, column=col, value=f); c.number_format = fmt
        c.font = F_LINK if f.startswith("=Inputs") else F_CALC
    ai.cell(row=r, column=4, value=how).font = F_NOTE
    for col in range(1, 5): ai.cell(row=r, column=col).border = BOX; ai.cell(row=r, column=col).alignment = WRAP
ai["B24"].fill = FILL_KEY
widths(ai, [44, 22, 22, 26, 20, 20])

# ---------------------------------------------------------------- One customer
oc = wb.create_sheet("One customer", 1)
title(oc, "Value for one customer, per year", "What the agent is worth to one mid-size buyer, and what it costs to run. Every number is a formula on the Inputs sheet.")
header(oc, 4, ["Line", "Cautious", "Likely", "Strong", "How it is calculated"])
lines = [
    ("flag", "Flagged supplier-side credit per year", lambda c: f"={ref('F',c)}*12", INR, "Flagged per month × 12", False),
    ("loss", "Credit no longer lost for good", lambda c: f"=B5_*({ref('L0',c)}-{ref('L1',c)})", INR, "Flagged × (share lost today − share lost with the agent)", True),
    ("extra", "Credit now fixed within the first cycle", lambda c: f"=B5_*({ref('FW',c)}-{ref('FT',c)})", INR, "Flagged × (share fixed in one cycle with the agent − today)", False),
    ("wc", "Value of receiving that credit sooner", lambda c: f"=B7_*{ref('CC',c)}*{ref('DS',c)}/365", INR, "Credit fixed sooner × cost of capital × days sooner ÷ 365", True),
    ("hrs", "GST team hours saved", lambda c: f"={ref('S',c)}*12*({ref('MT',c)}-{ref('MW',c)})/60", NUM, "Suppliers per month × 12 × minutes saved ÷ 60", False),
    ("time", "Value of time saved", lambda c: f"=B9_*{ref('H',c)}", INR, "Hours saved × loaded cost of an hour", True),
    ("tot", "Total value per customer per year", lambda c: "=B6_+B8_+B10_", INR, "Credit no longer lost + value of getting it sooner + value of time saved", False),
    ("totL", "Total value, in lakh", lambda c: "=B11_/100000", LAKH, "Total ÷ 1,00,000", False),
    ("lossShare", "Share of the value that comes from avoided loss", lambda c: "=IF(B11_=0,0,B6_/B11_)", PCT, "Credit no longer lost ÷ total value", False),
    ("price", "Price per customer per year", lambda c: f"={ref('P',c)}", INR, "From Inputs", False),
    ("ratio", "Value for each rupee of price", lambda c: "=IF(B14_=0,0,B11_/B14_)", X, "Total value ÷ price", False),
    ("payback", "Payback", lambda c: "=IF(B11_=0,0,B14_/(B11_/12))", MON, "Price ÷ monthly value", False),
    ("ai", "AI cost to run, per customer per year", lambda c: "='AI unit economics'!C26" if c == "D" else "='AI unit economics'!B26", INR, "Cautious uses the high-cost AI case; Likely and Strong use the likely case", False),
    ("aishare", "AI cost as a share of the price", lambda c: "=IF(B14_=0,0,B17_/B14_)", '0.00%', "AI cost ÷ price", False),
]
colmap = {"D": "B", "E": "C", "F": "D"}  # inputs column -> this sheet column
for i, (k, lab, fn, fmt, how, sub) in enumerate(lines):
    r = 5 + i
    oc.cell(row=r, column=1, value=lab).font = F_BOLD if k in ("tot", "ratio", "payback", "totL") else F_BASE
    for inc, (_, col) in zip(("B", "C", "D"), SC):
        f = fn(col).replace("B5_", f"{inc}5").replace("B6_", f"{inc}6").replace("B7_", f"{inc}7").replace("B8_", f"{inc}8").replace("B9_", f"{inc}9").replace("B10_", f"{inc}10").replace("B11_", f"{inc}11").replace("B14_", f"{inc}14").replace("B17_", f"{inc}17")
        c = oc[f"{inc}{r}"]; c.value = f; c.number_format = fmt
        c.font = F_LINK if ("Inputs!" in f or "AI unit" in f) and f.count("!") >= 1 and not any(op in f.replace("Inputs!", "") for op in "+-*/") and f.startswith("=Inputs") or f.startswith("='AI") else (F_CALCB if k in ("tot", "totL", "ratio", "payback") else F_CALC)
        c.border = BOX
        if k in ("tot", "totL"): c.fill = FILL_TOT
    oc.cell(row=r, column=5, value=how).font = F_NOTE
    oc.cell(row=r, column=1).border = BOX; oc.cell(row=r, column=5).border = BOX
oc["C11"].fill = FILL_KEY
oc["A20"] = "Reading it: avoided loss carries most of the value, so the loss rate today (hypothesis H4) decides the case. In the cautious case the agent is worth less than its price."; oc["A20"].font = F_NOTE
oc["A21"] = "Credit received sooner is valued only for the time gained (cost of capital), not counted as new money."; oc["A21"].font = F_NOTE
widths(oc, [46, 16, 16, 16, 70])
oc.freeze_panes = "B5"

# ---------------------------------------------------------------- Sensitivity
se = wb.create_sheet("Sensitivity", 2)
title(se, "What moves the value most", "Total value per customer per year with one input moved to its cautious or strong value and every other input at likely.")
header(se, 4, ["Input moved", "Value at cautious", "Value at strong", "Swing", "Cautious → strong value of the input"])
KEYS = ["F", "L0", "L1", "FT", "FW", "DS", "CC", "S", "MT", "MW", "H"]
def total_formula(over=None, col=None):
    g = lambda k: ref(k, col if (over == k) else "E")
    return (f"=12*{g('F')}*({g('L0')}-{g('L1')})"
            f"+12*{g('F')}*({g('FW')}-{g('FT')})*{g('CC')}*{g('DS')}/365"
            f"+{g('S')}*12*({g('MT')}-{g('MW')})/60*{g('H')}")
vals = {k: dict(zip(("D", "E", "F"), (lo, mid, hi))) for (k, _, _, lo, mid, hi, _, _) in INPUTS}
def total_py(over=None, col=None):
    v = lambda k: vals[k][col if over == k else "E"]
    return 12*v('F')*(v('L0')-v('L1')) + 12*v('F')*(v('FW')-v('FT'))*v('CC')*v('DS')/365 + v('S')*12*(v('MT')-v('MW'))/60*v('H')
order = sorted(KEYS, key=lambda k: -abs(total_py(k, "F") - total_py(k, "D")))
label = {k: l for (k, l, *_rest) in INPUTS}
fmts = {k: f for (k, _l, _u, _a, _b, _c, f, _n) in INPUTS}
for i, k in enumerate(order):
    r = 5 + i
    se.cell(row=r, column=1, value=label[k]).font = F_BOLD
    c = se.cell(row=r, column=2, value=total_formula(k, "D")); c.number_format = INR; c.font = F_CALC
    c = se.cell(row=r, column=3, value=total_formula(k, "F")); c.number_format = INR; c.font = F_CALC
    c = se.cell(row=r, column=4, value=f"=ABS(C{r}-B{r})"); c.number_format = INR; c.font = F_CALCB
    lo_s = f'TEXT({ref(k,"D")},"{ "0%" if fmts[k] in (PCT, PCT0) else "#,##0"}")'
    hi_s = f'TEXT({ref(k,"F")},"{ "0%" if fmts[k] in (PCT, PCT0) else "#,##0"}")'
    c = se.cell(row=r, column=5, value=f'={lo_s}&" → "&{hi_s}'); c.font = F_LINK
    for col in range(1, 6): se.cell(row=r, column=col).border = BOX
rl = 5 + len(order)
se.cell(row=rl + 1, column=1, value="Likely total, for reference").font = F_BOLD
c = se.cell(row=rl + 1, column=2, value=total_formula()); c.number_format = INR; c.font = F_CALCB; c.fill = FILL_TOT
se.cell(row=rl + 2, column=1, value="Rows are sorted by swing as built; if you change the inputs, re-sort by the Swing column.").font = F_NOTE
widths(se, [52, 18, 18, 18, 30])

# ---------------------------------------------------------------- Portfolio
pf = wb.create_sheet("Portfolio", 3)
title(pf, "Portfolio view", "If a number of customers adopt it: value created for them, revenue, and AI cost. Illustration only; adoption is an assumption.")
header(pf, 4, ["Line", "Cautious", "Likely", "Strong", "How it is calculated"])
pl = [
    ("Customers adopting", lambda c, oc_col: f"={ref('N', c)}", NUM, "From Inputs (assumption)"),
    ("Value created for customers per year", lambda c, oc_col: f"=B5_*'One customer'!{oc_col}11", INR, "Customers × value per customer"),
    ("Revenue per year", lambda c, oc_col: f"=B5_*'One customer'!{oc_col}14", INR, "Customers × price"),
    ("Revenue per year, in crore", lambda c, oc_col: "=B7_/10000000", '"₹"0.00" Cr"', "Revenue ÷ 1,00,00,000"),
    ("AI cost per year", lambda c, oc_col: f"=B5_*'One customer'!{oc_col}17", INR, "Customers × AI cost per customer"),
    ("AI cost as a share of revenue", lambda c, oc_col: "=IF(B7_=0,0,B9_/B7_)", '0.00%', "AI cost ÷ revenue"),
]
for i, (lab, fn, fmt, how) in enumerate(pl):
    r = 5 + i
    pf.cell(row=r, column=1, value=lab).font = F_BOLD
    for inc, (_, col) in zip(("B", "C", "D"), SC):
        f = fn(col, inc).replace("B5_", f"{inc}5").replace("B7_", f"{inc}7").replace("B9_", f"{inc}9")
        c = pf[f"{inc}{r}"]; c.value = f; c.number_format = fmt; c.border = BOX
        c.font = F_LINK if f.startswith("=Inputs") else F_CALC
    pf.cell(row=r, column=5, value=how).font = F_NOTE
    pf.cell(row=r, column=1).border = BOX; pf.cell(row=r, column=5).border = BOX
pf["A12"] = "Scenario columns compound: the cautious column has cautious adoption and cautious value per customer."; pf["A12"].font = F_NOTE
widths(pf, [40, 18, 18, 18, 44])

# ---------------------------------------------------------------- Demo check
dc = wb.create_sheet("Demo check")
title(dc, "Demo check: the prototype's simulated cycle", "Kestrel Auto Components (fictional), 15 Oct to 14 Nov 2026, as the prototype computes it. This sheet does not feed the model.")
header(dc, 4, ["On 14 Nov 2026", "Recommended steps taken on 15 Oct", "Nothing done", "Note"])
demo = [("Flagged", 2340396, 2340396, "29 invoices, 16 suppliers"),
        ("Recovered (in GSTR-2B)", 2236716, 642366, "Includes ₹52,200 already in the September GSTR-2B"),
        ("Lost for good (FY 2025-26, missed 30 Nov)", 0, 625500, "Nova ₹4,00,500, Vega ₹1,24,200, Rathi ₹1,00,800"),
        ("Still open", 90000, 1072530, "With action: GreenLeaf ₹86,400 promised for December; Krishna ₹3,600 broken promise"),
        ("Removed: duplicate in our own books", 13680, 0, "ST/2231-B, not real credit")]
for i, (lab, a, b, note) in enumerate(demo):
    r = 5 + i
    dc.cell(row=r, column=1, value=lab).font = F_BOLD
    for col, v in ((2, a), (3, b)):
        c = dc.cell(row=r, column=col, value=v); c.font = F_IN; c.number_format = INR; c.border = BOX
    dc.cell(row=r, column=4, value=note).font = F_NOTE
    dc.cell(row=r, column=1).border = BOX; dc.cell(row=r, column=4).border = BOX
dc["A11"] = "Check: parts add up to flagged"; dc["A11"].font = F_BOLD
dc["B11"] = "=IF(SUM(B6:B9)=B5,\"Yes\",\"No\")"; dc["C11"] = "=IF(SUM(C6:C9)=C5,\"Yes\",\"No\")"
dc["A12"] = "Real credit (flagged less the duplicate)"; dc["A12"].font = F_BOLD
dc["B12"] = "=B5-13680"; dc["C12"] = "=C5-13680"
dc["A13"] = "Share of real credit recovered"; dc["A13"].font = F_BOLD
dc["B13"] = "=B6/B12"; dc["C13"] = "=C6/C12"
dc["A14"] = "Share of real credit lost for good"; dc["A14"].font = F_BOLD
dc["B14"] = "=B7/B12"; dc["C14"] = "=C7/C12"
for a in ("B11", "C11"): dc[a].font = F_CALC
for a in ("B12", "C12"): dc[a].font = F_CALC; dc[a].number_format = INR
for a in ("B13", "C13", "B14", "C14"): dc[a].font = F_CALCB; dc[a].number_format = PCT
for r in range(11, 15):
    for col in range(1, 4): dc.cell(row=r, column=col).border = BOX
dc["A16"] = "The demo is built to show the mechanism, not the rate: it puts three FY 2025-26 cases into one month, so its loss without action (about 27%) is far above the 6% the model assumes."; dc["A16"].font = F_NOTE
dc["A17"] = "Blue figures are the prototype's outputs, copied as values. Source: the prototype's model (model.js) on fictional data."; dc["A17"].font = F_NOTE
widths(dc, [44, 30, 18, 80])

# ---------------------------------------------------------------- Inbox ranking
ib = wb.create_sheet("Inbox ranking")
title(ib, "How the inbox is ranked", "Protect what can't be undone first: last year's credit outranks bigger amounts that can still be claimed later.")
ib["A4"] = "Order of the rules"; ib["A4"].font = F_BOLD
steps = ["1. Group: Needs you, then Lost or missing, Promised, Waiting, Seen in IMS, Fixed.",
         "2. Within a group: suppliers with FY 2025-26 credit first (rule R8), because it cannot be claimed after 30 Nov 2026.",
         "3. Then, for promised suppliers, the earliest promised date first.",
         "4. Then the largest credit stuck first."]
for i, s in enumerate(steps): ib.cell(row=5 + i, column=1, value=s).font = F_BASE
header(ib, 10, ["Rank", "Supplier", "Status on 15 Oct 2026", "Last year's credit (₹)", "Credit stuck (₹)", "Why it sits here"])
RANK = [
    ("Nova Polymers", "Needs you", 400500, 537300, "Needs a decision; the most FY 2025-26 credit"),
    ("Vega Chemicals", "Needs you", 124200, 202500, "Needs a decision; FY 2025-26 credit"),
    ("Rathi Alloys", "Needs you", 100800, 158400, "Needs a decision; FY 2025-26 credit"),
    ("Patel Steel Traders", "Needs you", 0, 427500, "Needs a decision; largest current-year credit"),
    ("Delta Engineering Works", "Not contacted", 0, 226800, "Needs the first message"),
    ("GreenLeaf Facility Services", "Needs you", 0, 86400, "Contact has left"),
    ("Mehta Castings", "Fix in our books", 0, 59400, "Our entry is wrong"),
    ("Jyoti Industrial Gases", "Needs you", 0, 54900, "Angry denial"),
    ("Sai Tooling Solutions", "Needs you", 0, 38970, "Duplicate in our books"),
    ("Orbit Office Supplies", "Not contacted", 0, 3330, "Needs the first message"),
    ("Bharat Fasteners", "Needs you", 0, 2520, "Wants documents"),
    ("Kaveri Precision", "Promised 16 Oct", 0, 3060, "Earliest promise"),
    ("Sunrise Rubber Works", "Promised 11 Nov", 0, 78300, "Promise; larger credit first"),
    ("Shree Ganesh Packaging", "Promised 11 Nov", 0, 72216, "Promise"),
    ("Krishna Electricals", "Promised 11 Nov", 0, 3600, "Promise"),
    ("Apex Machine Tools", "Waiting for supplier", 0, 333000, "Out of office until 20 Oct"),
]
for i, (n, st, fy, op, why) in enumerate(RANK):
    r = 11 + i
    ib.cell(row=r, column=1, value=i + 1).font = F_BASE
    ib.cell(row=r, column=2, value=n).font = F_BOLD
    ib.cell(row=r, column=3, value=st).font = F_BASE
    c = ib.cell(row=r, column=4, value=fy); c.font = F_IN; c.number_format = INR
    c = ib.cell(row=r, column=5, value=op); c.font = F_IN; c.number_format = INR
    ib.cell(row=r, column=6, value=why).font = F_NOTE
    for col in range(1, 7): ib.cell(row=r, column=col).border = BOX
rr = 11 + len(RANK)
ib.cell(row=rr, column=2, value="Total").font = F_BOLD
c = ib.cell(row=rr, column=4, value=f"=SUM(D11:D{rr-1})"); c.font = F_CALCB; c.number_format = INR
c = ib.cell(row=rr, column=5, value=f"=SUM(E11:E{rr-1})"); c.font = F_CALCB; c.number_format = INR
ib.cell(row=rr + 1, column=2, value="Needs you (first 11 rows)").font = F_BOLD
c = ib.cell(row=rr + 1, column=5, value="=SUM(E11:E21)"); c.font = F_CALCB; c.number_format = INR
ib.cell(row=rr + 2, column=2, value="Order and amounts as the prototype shows them on 15 Oct 2026 (fictional data). Credit stuck excludes ₹52,200 already recovered.").font = F_NOTE
widths(ib, [7, 30, 22, 20, 18, 46])

# ---------------------------------------------------------------- Read me
rm = wb.create_sheet("Read me", 0)
title(rm, "Supplier Follow-up Agent: impact model", "Illustrative. Every input is an assumption until shadow mode measures it. Designed and directed by Rohit Kumar; built with Claude; 11 Oct 2026.")
rm["A4"] = "The answer, in the likely case"; rm["A4"].font = F_BOLD; rm["A4"].fill = FILL_SEC
key = [
    ("Value for one customer per year", "='One customer'!C11", INR),
    ("Value for each rupee of price", "='One customer'!C15", X),
    ("Payback", "='One customer'!C16", MON),
    ("Share of the value from credit no longer lost", "='One customer'!C13", PCT),
    ("Value in the cautious case (below the ₹2 L price)", "='One customer'!B11", INR),
    ("AI cost per supplier per cycle", "='AI unit economics'!B24", INR2),
    ("AI cost as a share of the price", "='One customer'!C18", '0.00%'),
]
for i, (lab, f, fmt) in enumerate(key):
    r = 5 + i
    rm.cell(row=r, column=1, value=lab).font = F_BASE
    c = rm.cell(row=r, column=2, value=f); c.font = F_LINKB; c.number_format = fmt
    rm.cell(row=r, column=1).border = BOX; c.border = BOX
rm["A13"] = "What it means"; rm["A13"].font = F_BOLD; rm["A13"].fill = FILL_SEC
means = ["Avoided loss carries about three quarters of the value, so the loss rate today (hypothesis H4) decides the business case.",
         "With every input at its cautious end at once, the agent is worth about ₹1.8 L a year, less than a ₹2 L price. No single input on its own pulls the likely case below ₹6.8 L.",
         "AI cost is not the constraint: well under 1% of the price even at older model prices.",
         "So the first two weeks of the rollout measure the loss rate before anything is built."]
for i, s in enumerate(means): rm.cell(row=14 + i, column=1, value=s).font = F_BASE
rm["A19"] = "Sheets"; rm["A19"].font = F_BOLD; rm["A19"].fill = FILL_SEC
guide = [("Inputs", "Every assumption in three scenarios, with the reason and how it gets measured. Edit the blue cells."),
         ("One customer", "Value per customer per year, price, payback and AI cost."),
         ("Sensitivity", "Which input moves the value most, one at a time."),
         ("Portfolio", "Value, revenue and AI cost if a number of customers adopt."),
         ("AI unit economics", "Tokens per job, list prices and cost per supplier per cycle."),
         ("Demo check", "The prototype's simulated cycle, with and without the recommended steps."),
         ("Inbox ranking", "How the inbox orders suppliers, on the demo data.")]
for i, (s, d) in enumerate(guide):
    rm.cell(row=20 + i, column=1, value=s).font = F_BOLD
    rm.cell(row=20 + i, column=2, value=d).font = F_BASE
rm["A28"] = "Colours"; rm["A28"].font = F_BOLD; rm["A28"].fill = FILL_SEC
rm["A29"] = "Blue = typed input · black = formula · green = taken from another sheet · yellow = the likely case quoted in the document and slides."; rm["A29"].font = F_BASE
rm["A31"] = "Sources"; rm["A31"].font = F_BOLD; rm["A31"].fill = FILL_SEC
src = ["Model prices: Anthropic pricing page, platform.claude.com/docs/en/about-claude/pricing (opened 11 Oct 2026).",
       "Exchange rate: Upstox currency converter, ₹96.64 per US$ on 11 Oct 2026.",
       "Deadline for FY 2025-26 credit: Section 16(4) of the CGST Act, 30 Nov 2026. Reversal after 180 days unpaid: Rule 37. Supplier not filing GSTR-3B: Rule 37A.",
       "Demo figures: the prototype, on fictional data. No customer data was used anywhere in this model."]
for i, s in enumerate(src): rm.cell(row=32 + i, column=1, value=s).font = F_NOTE
widths(rm, [58, 64])

from openpyxl.worksheet.properties import PageSetupProperties
for ws in wb.worksheets:
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.tabColor = "142338" if ws.title in ("Read me",) else ("405AE8" if ws.title == "Inputs" else "DFE5EE")
wb.save(OUT)
print("saved", OUT)
for k in order:
    print(f"{label[k]:55s} cautious {total_py(k,'D'):>12,.0f}  strong {total_py(k,'F'):>12,.0f}  swing {abs(total_py(k,'F')-total_py(k,'D')):>12,.0f}")
print("likely", total_py())
